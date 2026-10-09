import csv
import io
import zipfile
from datetime import date, timedelta

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Max
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    BoardRevision, Classroom, DailyBoard, DailySubmission, PraiseItem,
    RosterImport, Student, TeacherAssignment,
)
from .permissions import can_manage_class, is_admin


def parse_day(value):
    try:
        result = date.fromisoformat(value) if value else timezone.localdate()
    except (ValueError, TypeError):
        raise ValidationError("日期格式无效。")
    if result > timezone.localdate():
        raise ValidationError("不能填写未来日期。")
    return result


def require_version(board, value):
    try:
        valid = int(value) == board.version
    except (TypeError, ValueError):
        valid = False
    if not valid:
        raise ValidationError("记录已被其他人更新，请刷新页面后再操作。")


def bump_version(board):
    board.version += 1
    board.save(update_fields=("version", "updated_at"))


@transaction.atomic
def save_items(user, board_id, assignment_id, items, version, submitted=True):
    board = DailyBoard.objects.select_for_update().select_related("classroom").get(pk=board_id)
    require_version(board, version)
    if board.is_finalized:
        raise ValidationError("榜单已定稿，请联系班主任重新打开。")
    manager = can_manage_class(user, board.classroom)
    assignment = TeacherAssignment.objects.select_related("teacher", "subject").get(pk=assignment_id)
    if assignment.classroom_id != board.classroom_id:
        raise PermissionDenied
    if not manager:
        if (
            assignment.teacher_id != user.id or not assignment.is_active
            or not assignment.subject.is_active or not assignment.classroom.is_active
            or board.day != timezone.localdate()
        ):
            raise PermissionDenied
    if not isinstance(items, list) or not 1 <= len(items) <= 30:
        raise ValidationError("每次填报须包含 1 至 30 个表扬项目。")
    roster = {student.pk: student for student in board.classroom.students.filter(is_active=True)}
    cleaned = []
    for index, item in enumerate(items):
        if not isinstance(item, dict):
            raise ValidationError("填报格式无效。")
        title = str(item.get("title", "")).strip()
        note = str(item.get("note", "")).strip()
        if not title or len(title) > 120 or len(note) > 400:
            raise ValidationError("项目标题须为 1 至 120 字，说明不能超过 400 字。")
        all_students = item.get("all_students") is True
        try:
            student_ids = set(int(value) for value in item.get("student_ids", []))
        except (TypeError, ValueError):
            raise ValidationError("所选学生无效。")
        if all_students:
            students = list(roster.values())
        else:
            if not student_ids or not student_ids.issubset(roster):
                raise ValidationError("请选择本班在读学生。")
            students = [student for student in roster.values() if student.pk in student_ids]
        if not students:
            raise ValidationError("班级尚无在读学生，请先维护名单。")
        cleaned.append(dict(
            title=title, note=note, all_students=all_students, sort_order=index,
            student_snapshots=[{"id": s.pk, "student_no": s.student_no, "name": s.name} for s in students],
        ))
    submission, _ = DailySubmission.objects.get_or_create(board=board, assignment=assignment)
    submission.items.all().delete()
    PraiseItem.objects.bulk_create([PraiseItem(submission=submission, **values) for values in cleaned])
    submission.is_submitted = submitted
    submission.submitted_at = timezone.now() if submitted else None
    submission.save()
    bump_version(board)
    return submission


def board_snapshot(board):
    sections = {}
    submissions = board.submissions.filter(is_submitted=True).select_related(
        "assignment__subject", "assignment__teacher"
    ).prefetch_related("items")
    for submission in submissions:
        subject = submission.assignment.subject
        section = sections.setdefault(subject.pk, {"subject": subject.name, "items": []})
        for item in submission.items.all():
            section["items"].append({
                "id": item.pk, "title": item.title, "note": item.note,
                "all_students": item.all_students, "students": item.student_snapshots,
                "teacher": str(submission.assignment.teacher), "sort_order": item.sort_order,
            })
    for section in sections.values():
        section["items"].sort(key=lambda item: (item["sort_order"], item["id"]))
    return {
        "classroom": board.classroom.name, "day": board.day.isoformat(),
        "title": board.title or "今日表扬榜", "footer": board.footer,
        "sections": list(sections.values()),
    }


@transaction.atomic
def change_board_state(user, board_id, action, version, title="", footer=""):
    board = DailyBoard.objects.select_for_update().select_related("classroom").get(pk=board_id)
    if not can_manage_class(user, board.classroom):
        raise PermissionDenied
    require_version(board, version)
    if action == "reopen":
        if not board.is_finalized:
            raise ValidationError("榜单已经处于进行中。")
        board.state = DailyBoard.State.OPEN
        board.finalized_at = None
        board.finalized_by = None
    else:
        if board.is_finalized:
            raise ValidationError("榜单已定稿。")
        if len(title) > 120 or len(footer) > 160:
            raise ValidationError("标题或落款过长。")
        board.title, board.footer = title.strip(), footer.strip()
        if action == "finalize":
            snapshot = board_snapshot(board)
            if not any(section["items"] for section in snapshot["sections"]):
                raise ValidationError("没有已提交的表扬项目，不能定稿。")
            number = (board.revisions.aggregate(number=Max("number"))["number"] or 0) + 1
            BoardRevision.objects.create(board=board, number=number, snapshot=snapshot, created_by=user)
            board.state = DailyBoard.State.FINALIZED
            board.finalized_at = timezone.now()
            board.finalized_by = user
        elif action != "metadata":
            raise ValidationError("无效操作。")
    board.version += 1
    board.save()
    return board


@transaction.atomic
def delete_item(user, item_id, version):
    item = PraiseItem.objects.select_related("submission__board__classroom").get(pk=item_id)
    board = DailyBoard.objects.select_for_update().get(pk=item.submission.board_id)
    if not can_manage_class(user, item.submission.board.classroom):
        raise PermissionDenied
    require_version(board, version)
    if board.is_finalized:
        raise ValidationError("榜单已定稿。")
    item.delete()
    bump_version(board)


@transaction.atomic
def reorder_item(user, item_id, direction, version):
    item = PraiseItem.objects.select_related("submission__board__classroom").get(pk=item_id)
    board = DailyBoard.objects.select_for_update().get(pk=item.submission.board_id)
    if not can_manage_class(user, item.submission.board.classroom):
        raise PermissionDenied
    require_version(board, version)
    if board.is_finalized or direction not in ("up", "down"):
        raise ValidationError("不能调整项目顺序。")
    items = list(PraiseItem.objects.filter(
        submission__board=board,
        submission__is_submitted=True,
        submission__assignment__subject=item.submission.assignment.subject_id,
    ).order_by("sort_order", "id"))
    position = next(i for i, other in enumerate(items) if other.pk == item.pk)
    target = position + (-1 if direction == "up" else 1)
    if 0 <= target < len(items):
        items[position], items[target] = items[target], items[position]
        for order, other in enumerate(items):
            other.sort_order = order
        PraiseItem.objects.bulk_update(items, ["sort_order"])
        bump_version(board)


def read_roster(upload):
    if upload.size > 2 * 1024 * 1024:
        raise ValidationError("文件不能超过 2 MB。")
    extension = upload.name.lower().rsplit(".", 1)[-1]
    try:
        if extension == "csv":
            content = upload.read().decode("utf-8-sig")
            raw = list(csv.reader(io.StringIO(content)))
        elif extension == "xlsx":
            with zipfile.ZipFile(upload) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValidationError("Excel 文件解压后过大，请只上传学生名单。")
            upload.seek(0)
            workbook = load_workbook(upload, read_only=True, data_only=False)
            sheet = workbook.active
            if sheet.max_row > 2001 or sheet.max_column > 20:
                raise ValidationError("名单文件最多支持 2000 名学生和 20 列。")
            raw = list(sheet.iter_rows(values_only=True))
            workbook.close()
        else:
            raise ValidationError("仅支持 UTF-8 CSV 或 XLSX 文件。")
    except (UnicodeError, csv.Error):
        raise ValidationError("CSV 文件编码无效，请使用 UTF-8 保存。")
    except ValidationError:
        raise
    except Exception:
        raise ValidationError("无法读取文件，请检查文件内容和格式。")
    if not raw or len(raw) > 2001:
        raise ValidationError("名单文件须包含表头，且最多支持 2000 名学生。")
    headers = [str(value or "").strip() for value in raw[0]]
    if headers.count("学号") != 1 or headers.count("姓名") != 1:
        raise ValidationError("表头须包含唯一的“学号”和“姓名”列。")
    number_col, name_col = headers.index("学号"), headers.index("姓名")
    rows, errors, seen = [], [], set()
    for line, record in enumerate(raw[1:], 2):
        if not any(value not in (None, "") for value in record):
            continue
        number = str(record[number_col] if number_col < len(record) and record[number_col] is not None else "").strip()
        name = str(record[name_col] if name_col < len(record) and record[name_col] is not None else "").strip()
        if not number or len(number) > 40 or not name or len(name) > 80 or number.startswith(("=", "+", "-", "@")) or name.startswith("="):
            errors.append(f"第 {line} 行：学号或姓名为空、过长，或包含公式。")
        if number in seen:
            errors.append(f"第 {line} 行：学号 {number} 在文件中重复。")
        seen.add(number)
        rows.append({"student_no": number, "name": name})
    if not rows:
        errors.append("文件中没有学生记录。")
    return rows, errors


@transaction.atomic
def confirm_import(user, import_id):
    preview = RosterImport.objects.select_for_update().select_related("classroom").get(pk=import_id)
    if not can_manage_class(user, preview.classroom) or (not is_admin(user) and preview.created_by_id != user.pk):
        raise PermissionDenied
    Classroom.objects.select_for_update().get(pk=preview.classroom_id)
    if preview.errors or preview.state != RosterImport.State.PREVIEW:
        raise ValidationError("该批名单无效或已导入。")
    if preview.created_at < timezone.now() - timedelta(hours=24):
        raise ValidationError("预览已过期，请重新上传。")
    for row in preview.rows:
        Student.objects.update_or_create(
            classroom=preview.classroom, student_no=row["student_no"],
            defaults={"name": row["name"], "is_active": True},
        )
    preview.state = RosterImport.State.IMPORTED
    preview.save(update_fields=("state",))
    return len(preview.rows)
